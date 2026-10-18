"""Exact source selection, fresh grants, source drift and local calculation."""

import hashlib
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from pydantic import ValidationError

from backend import auth
from backend.dependencies import store
from backend.models import (
    AccountCreate,
    BookingCreate,
    ContactCreate,
    DocumentCreate,
    PortfolioCreate,
    PropertyCreate,
    TenantCreate,
    UnitCreate,
)
from backend.services import assistant_context as context
from backend.services import file_storage


@pytest.fixture
def estate(monkeypatch, tmp_path):
    auth.clear_users()
    store.clear_all()
    store.__dict__.pop("_personal_workspace", None)
    monkeypatch.setattr(file_storage, "_storage", file_storage.LocalStorage(tmp_path / "uploads"))
    north = store.create_portfolio(PortfolioCreate(name="Nord"))
    south = store.create_portfolio(PortfolioCreate(name="Süd"))
    prop = store.create_property(PropertyCreate(portfolio_id=north.id, name="Nordhaus", property_type="residential"))
    foreign = store.create_property(PropertyCreate(portfolio_id=south.id, name="Südhaus", property_type="residential"))
    unit = store.create_unit(UnitCreate(property_id=prop.id, label="Links", unit_type="apartment"))
    tenant = store.create_tenant(TenantCreate(full_name="Ausgewählte Partei"))
    contact = store.create_contact(ContactCreate(first_name="Ada", iban="HIDDEN-BANK-DATA", tax_id="HIDDEN-TAX"))
    user = auth.register_user("assistant-user", "assistant@example.test", "Test", "Secret123", "eigentuemer")
    yield SimpleNamespace(north=north, south=south, prop=prop, foreign=foreign, unit=unit,
                          tenant=tenant, contact=contact, user=user)
    auth.clear_users()
    store.clear_all()
    store.__dict__.pop("_personal_workspace", None)


def document(estate, text="Lieferant: Testfirma\nRechnungsnr: R-2026\nRechnungsdatum: 18.10.2026\nGesamtbetrag: 119,00"):
    from reportlab.pdfgen.canvas import Canvas
    output = BytesIO()
    pdf = Canvas(output)
    for index, line in enumerate(text.splitlines()):
        pdf.drawString(50, 780 - index * 20, line)
    pdf.save()
    content = output.getvalue()
    file_storage.get_file_storage().save("documents/assistant.pdf", BytesIO(content))
    row = store.create_document(DocumentCreate(property_id=estate.prop.id, title="Originalrechnung",
                                                file_url="/uploads/documents/assistant.pdf"))
    return row, content


def test_only_explicit_sources_are_disclosed_and_bank_data_is_absent(estate):
    bundle = context.build_context(store, {"property_ids": [estate.prop.id, estate.prop.id],
                                           "contact_ids": [estate.contact.id]}, estate.user.id)
    assert {(s["kind"], s["id"]) for s in bundle["sources"]} == {
        ("property", estate.prop.id), ("contact", estate.contact.id)}
    assert len(bundle["sha256"]) == 64
    assert all(s["revision"] and s["href"].startswith("/") for s in bundle["sources"])
    assert "HIDDEN" not in str(bundle)
    assert all(s["content_role"] == "untrusted_evidence" for s in bundle["sources"])
    assert context.revalidate_context(store, bundle, estate.user.id) == bundle


def test_foreign_exact_id_is_404_and_revocation_applies_to_existing_bundle(estate):
    reader = auth.register_user("scoped", "scoped@example.test", "Scoped", "Secret123", "readonly",
                                portfolio_access="selected", portfolio_ids=[estate.north.id])
    with pytest.raises(HTTPException) as denied:
        context.build_context(store, {"property_ids": [estate.foreign.id]}, reader.id)
    assert denied.value.status_code == 404
    bundle = context.build_context(store, {"property_ids": [estate.prop.id]}, reader.id)
    auth.update_user(reader.id, {"portfolio_access": "selected", "portfolio_ids": []})
    with pytest.raises(HTTPException) as revoked:
        context.assert_context_access(store, bundle, reader.id)
    assert revoked.value.status_code == 404


def test_normal_edit_keeps_old_reply_readable_but_prevents_stale_preview(estate):
    bundle = context.build_context(store, {"property_ids": [estate.prop.id]}, estate.user.id)
    store.update_property(estate.prop.id, PropertyCreate(portfolio_id=estate.north.id, name="Geändert", property_type="residential"))
    context.assert_context_access(store, bundle, estate.user.id)
    with pytest.raises(HTTPException) as changed:
        context.revalidate_context(store, bundle, estate.user.id)
    assert changed.value.status_code == 409


def test_caller_cannot_modify_context_payload_behind_hash(estate):
    bundle = context.build_context(store, {"property_ids": [estate.prop.id]}, estate.user.id)
    bundle["sources"][0]["values"]["name"] = "Falsche Behauptung"
    with pytest.raises(HTTPException) as changed:
        context.revalidate_context(store, bundle, estate.user.id)
    assert changed.value.status_code == 409


@pytest.mark.parametrize("payload", [{"all": True}, {"property_ids": [" "]}, {"unit_ids": [" a"]}])
def test_selection_rejects_implicit_or_invalid_sources(payload):
    with pytest.raises(ValidationError):
        context.ContextSelection.model_validate(payload)


def test_document_bytes_are_part_of_source_hash_even_without_metadata_edit(estate):
    row, _ = document(estate)
    bundle = context.build_context(store, {"document_ids": [row.id]}, estate.user.id)
    assert "file_url" not in bundle["sources"][0]["values"]
    storage = file_storage.get_file_storage()
    assert isinstance(storage, file_storage.LocalStorage)
    # Simulate disk corruption behind the immutable storage API, in tmp_path.
    storage._path("documents/assistant.pdf").write_bytes(b"changed")
    with pytest.raises(HTTPException) as changed:
        context.revalidate_context(store, bundle, estate.user.id)
    assert changed.value.status_code == 409


def test_local_document_extraction_uses_real_pdf_regex_without_model(estate, monkeypatch):
    from backend.services.ai.hf_runtime import runtime
    monkeypatch.setattr(runtime, "get_pipeline", lambda *a, **k: pytest.fail("No model pipeline may load"))
    monkeypatch.setattr(runtime, "get_embedder", lambda *a, **k: pytest.fail("No embedding model may load"))
    row, _ = document(estate)
    result = context.local_analysis(store, {"document_ids": [row.id], "unit_ids": [estate.unit.id]},
                                    estate.user.id, extract_documents=True)
    extracted = result["documents"][0]
    assert extracted["origin"] == {"method": "regex", "model": None}
    assert extracted["invoice_number"] == "R-2026"
    assert extracted["total_amount"] == 119.0
    assert "R-2026" in extracted["text"] and not extracted["text_truncated"] and extracted["omissions"] == []
    assert extracted["source"]["id"] == row.id
    assert "file_url" not in extracted["source"]
    assert result["findings"][0]["code"] == "unit_area_missing"
    assert result["findings"][0]["evidence"][0]["id"] == estate.unit.id
    assert not result["model_used"] and not result["domain_writes"]


def test_local_document_text_is_bounded_with_full_text_hash_and_one_extraction(estate, monkeypatch):
    from backend.services import ocr_service
    row, _ = document(estate)
    full_text = "A" * 64_000 + "UNSEEN-CONTENT"
    calls = []

    def extract(content, extension):
        calls.append(extension)
        return full_text

    monkeypatch.setattr(ocr_service, "extract_text_from_bytes", extract)
    result = context.local_analysis(store, {"document_ids": [row.id]}, estate.user.id, extract_documents=True)
    extracted = result["documents"][0]
    assert len(extracted["text"]) == 64_000 and "UNSEEN-CONTENT" not in extracted["text"]
    assert extracted["text_truncated"] and extracted["omissions"]
    assert extracted["source"]["text_sha256"] == hashlib.sha256(full_text.encode()).hexdigest()
    assert calls == ["pdf"]


def test_local_finance_uses_exact_engine_with_explicit_property_and_source_links(estate):
    account = store.create_account(AccountCreate(portfolio_id=estate.north.id, name="Konto", account_type="bank"))
    first = store.create_booking(BookingCreate(account_id=account.id, property_id=estate.prop.id,
                                               booking_date=date(2026, 10, 1), amount=0.1))
    store.create_booking(BookingCreate(account_id=account.id, property_id=estate.prop.id,
                                      booking_date=date(2026, 10, 2), amount=0.2))
    store.create_booking(BookingCreate(account_id=account.id, property_id=estate.foreign.id,
                                      booking_date=date(2026, 10, 2), amount=999))
    definition = {"kind": "cash-v2", "as_of": "2026-10-18", "date_from": "2026-10-01", "property_id": estate.prop.id}
    result = context.local_analysis(store, {"property_ids": [estate.prop.id]}, estate.user.id, report=definition)
    assert result["financial_engine"] == "finance_reports.compute"
    assert result["finance"]["totals"]["net"] == "0.30"
    assert any(s["identity"] == "booking:" + first.id and s["url"].startswith("/bookings?")
               for s in result["finance"]["sources"])
    with pytest.raises(HTTPException) as implicit:
        context.local_analysis(store, {}, estate.user.id, report=definition)
    assert implicit.value.status_code == 422


def test_local_statistics_work_without_model_or_selected_sources(estate):
    result = context.local_analysis(store, {}, estate.user.id)
    assert result["mode"] == "local" and result["finance"] is None
    assert result["context"]["sources"] == []


@pytest.mark.parametrize("extension", ["docx", "txt", "csv"])
def test_existing_documents_use_same_local_text_reader_as_chat_uploads(estate, extension):
    contents = "Vorhandener Vertragsbeleg mit eindeutiger Referenz V-2026"
    if extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph(contents)
        doc.sections[0].header.paragraphs[0].text = "Kopfzeile geprüft"
        stream = BytesIO()
        doc.save(stream)
        raw = stream.getvalue()
    else:
        raw = contents.encode("utf-8")
    key = "documents/existing." + extension
    file_storage.get_file_storage().save(key, BytesIO(raw))
    row = store.create_document(DocumentCreate(property_id=estate.prop.id, title="Vorhandener Beleg", file_url="/uploads/" + key))
    result = context.local_analysis(store, {"document_ids": [row.id]}, estate.user.id, extract_documents=True)
    extracted = result["documents"][0]
    assert contents in extracted["text"] and not extracted["text_truncated"]
    assert extracted["source"]["file_sha256"] == hashlib.sha256(raw).hexdigest()
    if extension == "docx":
        assert "Kopfzeile geprüft" in extracted["text"]
